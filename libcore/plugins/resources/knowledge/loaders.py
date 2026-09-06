"""文档加载器（原生重写，不 import 旧 app.runtime）。

对齐旧 ``TextLoader`` / ``DocxLoader`` / ``PdfLoader``：
- 按扩展名路由到对应加载器；
- 可选依赖（python-docx / pypdf）缺失时降级为"不支持该格式"，不抛致命错误（无感降级）。
"""
from __future__ import annotations

from typing import Dict, List, Optional

from .documents import KBDocument


class DocumentLoader:
    """按扩展名路由的文档加载器。"""

    _TEXT_EXTENSIONS = {
        "txt", "md", "markdown", "rst", "log",
        "csv", "json", "yaml", "yml", "toml", "ini", "cfg",
        "py", "js", "ts", "tsx", "jsx", "html", "htm", "css", "scss",
        "java", "c", "cpp", "h", "hpp", "cs", "go", "rs", "rb", "php",
        "sh", "bash", "zsh", "fish", "sql",
        "gitignore", "dockerfile", "env",
    }

    def load(self, file_path: str) -> List[KBDocument]:
        """按扩展名路由加载文档。未知/不支持格式 → 返回空列表（不抛致命错误）。"""
        ext = file_path.rsplit(".", 1)[-1].lower() if "." in file_path else ""
        if ext in self._TEXT_EXTENSIONS:
            return self._load_text(file_path)
        if ext == "docx":
            return self._load_docx(file_path)
        if ext == "pdf":
            return self._load_pdf(file_path)
        return []

    def _load_text(self, file_path: str) -> List[KBDocument]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return [KBDocument(text=text, metadata={"source": file_path})]

    def _load_docx(self, file_path: str) -> List[KBDocument]:
        try:
            from docx import Document
        except ImportError:
            return []
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        text = "\n\n".join([p for p in paragraphs if p and p.strip()])
        return [KBDocument(text=text, metadata={"source": file_path})]

    def _load_pdf(self, file_path: str) -> List[KBDocument]:
        try:
            from pypdf import PdfReader
        except ImportError:
            return []
        reader = PdfReader(file_path)
        results: List[KBDocument] = []
        for idx, page in enumerate(reader.pages):
            try:
                text = (page.extract_text() or "").strip()
            except Exception:
                text = ""
            if not text:
                continue
            results.append(KBDocument(
                text=text,
                metadata={"source": file_path, "page": idx + 1,
                          "total_pages": len(reader.pages)},
            ))
        return results

    def supported_extensions(self) -> List[str]:
        return sorted(self._TEXT_EXTENSIONS | {"docx", "pdf"})


def load_document(file_path: str) -> List[KBDocument]:
    """便捷入口：按扩展名加载文档。"""
    return DocumentLoader().load(file_path)
